"""Helpers for document tests: small PDFs and Word files generated on the fly.

All content is fictitious. PDFs use a core font, so text is limited to Latin-1
(å, ä and ö work; typographic quotes and dashes are tested on plain strings instead).
"""

from __future__ import annotations

import io
from collections.abc import Callable, Sequence

import docx
import pytest
from fpdf import FPDF

MakeFile = Callable[..., bytes]


def build_pdf(pages: Sequence[str]) -> bytes:
    """One PDF page per entry; explicit newlines become line breaks."""

    pdf = FPDF()
    pdf.set_auto_page_break(auto=False)
    pdf.set_font("Helvetica", size=11)
    for text in pages:
        pdf.add_page()
        if text:
            pdf.multi_cell(0, 6, text)
    return bytes(pdf.output())


def build_docx(paragraphs: Sequence[str], table: Sequence[Sequence[str]] = ()) -> bytes:
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row, values in zip(grid.rows, table, strict=True):
            for cell, value in zip(row.cells, values, strict=True):
                cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def make_pdf() -> MakeFile:
    return build_pdf


@pytest.fixture
def make_docx() -> MakeFile:
    return build_docx


TENDER_PAGES: tuple[str, ...] = (
    "FIKTIVT UNDERLAG\n1 Administrativa föreskrifter\nSista anbudsdag är 2026-11-14. "
    "Frågor ställs i e-Avrop senast 2026-11-04.",
    "2 Kvalificering\nAnbudsgivaren ska inneha en giltig ansvarsförsäk-\nring om minst 5 000 000 kronor "
    "per skadetillfälle.\nFörsäkringsbevis ska bifogas anbudet.",
    "3 Kravspecifikation\nArbetet ska utföras enligt AMA Hus 21. Målning ska ske med\n"
    "miljömärkt färg i samråd med förskolans personal.",
    "4 Utvärdering\nPris 70 %, kvalitet 30 %. Anbudet ska vara giltigt i 90 dagar efter sista anbudsdag.",
)


@pytest.fixture
def tender_pages() -> list[str]:
    return list(TENDER_PAGES)

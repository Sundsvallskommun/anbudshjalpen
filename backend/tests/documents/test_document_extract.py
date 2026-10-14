"""Text extraction per page for PDF, DOCX, text and markdown."""

from __future__ import annotations

import io
import os
import zipfile
from collections.abc import Callable

import docx
import pytest
from fpdf import FPDF

from anbudsklar.documents import extract
from anbudsklar.documents.extract import (
    DOCX_MIMETYPE,
    MAX_CHARS,
    MAX_PAGES,
    PDF_MIMETYPE,
    PSEUDO_PAGE_CHARS,
    TEXT_MIMETYPE,
    DocumentTooLarge,
    UnreadableDocument,
    UnsupportedDocument,
    clean_page_text,
    extract_pages,
    paginate,
    resolve_mimetype,
)

MakeFile = Callable[..., bytes]

# --- clean_page_text ------------------------------------------------------------------


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("Kvalifi-\nceringskraven gäller", "Kvalificeringskraven gäller"),
        ("kommunens IT-\nsystem", "kommunens IT-system"),
        ("via e-\npost", "via e-post"),
        ("bygg-\noch anläggningsarbeten", "bygg- och anläggningsarbeten"),
        ("vecka 2026-\n10-14", "vecka 2026-10-14"),
        ("anbuds\N{SOFT HYPHEN}\ngivare och för\N{SOFT HYPHEN}skola", "anbudsgivare och förskola"),
        ("\N{LATIN SMALL LIGATURE FI}rma", "firma"),
        ("5\N{NO-BREAK SPACE}000\N{NARROW NO-BREAK SPACE}kr", "5 000 kr"),
        ("a\N{ZERO WIDTH SPACE}b", "ab"),
        ("  Rad ett   med  mellanrum  \r\nRad två\t\tslut ", "Rad ett med mellanrum\nRad två slut"),
        ("Stycke ett\n\n\n\n\nStycke två", "Stycke ett\n\nStycke två"),
        ("Fo\N{COMBINING DIAERESIS}rskola", "Förskola"),
        ("Pris \N{EN DASH} 70 %", "Pris \N{EN DASH} 70 %"),
    ],
)
def test_clean_page_text(raw: str, expected: str) -> None:
    assert clean_page_text(raw) == expected


def test_clean_page_text_keeps_list_lines() -> None:
    text = "Anbudet ska innehålla:\n- prisbilaga\n- referensblankett"
    assert clean_page_text(text) == text


# --- pseudo-pages ---------------------------------------------------------------------


def test_paginate_groups_paragraphs_without_splitting_them() -> None:
    paragraphs = [f"Stycke {n}. " + "ord " * 150 for n in range(12)]
    pages = paginate(paragraphs, separator="\n\n", size=2_000)
    assert len(pages) > 1
    assert all(len(page) <= 2_000 for page in pages)
    assert "\n\n".join(pages).split("\n\n") == [p.strip() for p in paragraphs]


def test_paginate_splits_an_oversized_paragraph_at_sentence_ends() -> None:
    paragraph = " ".join(f"Mening nummer {n} om kraven i upphandlingen." for n in range(200))
    pages = paginate([paragraph], separator="\n", size=1_000)
    assert len(pages) > 1
    assert all(len(page) <= 1_000 for page in pages)
    assert all(page.endswith(".") for page in pages)


def test_paginate_cuts_a_block_without_spaces_into_full_pages() -> None:
    pages = paginate(["x" * 2_500], separator="\n", size=1_000)
    assert [len(page) for page in pages] == [1_000, 1_000, 500]


# --- formats --------------------------------------------------------------------------


def test_pdf_keeps_physical_pages(make_pdf: MakeFile) -> None:
    data = make_pdf(
        ["Sida ett om förskolan i Matfors.", "Sida två: kvalificering.", "Sida tre: utvärdering."]
    )
    pages = extract_pages(data, PDF_MIMETYPE, "underlag.pdf")
    assert pages == ["Sida ett om förskolan i Matfors.", "Sida två: kvalificering.", "Sida tre: utvärdering."]


def test_pdf_rejoins_words_hyphenated_over_line_breaks(make_pdf: MakeFile) -> None:
    data = make_pdf(
        ["Kvalifi-\nceringskraven gäller alla.\nKommunens IT-\nsystem och e-\npost, bygg-\noch måleri."]
    )
    [page] = extract_pages(data, PDF_MIMETYPE, "underlag.pdf")
    assert page == "Kvalificeringskraven gäller alla.\nKommunens IT-system och e-post, bygg- och måleri."


def test_pdf_blank_page_keeps_its_number(make_pdf: MakeFile) -> None:
    pages = extract_pages(make_pdf(["Första sidan.", "", "Tredje sidan."]), PDF_MIMETYPE, "underlag.pdf")
    assert pages == ["Första sidan.", "", "Tredje sidan."]


def test_pdf_is_detected_without_a_useful_mimetype(make_pdf: MakeFile) -> None:
    data = make_pdf(["Innehåll."])
    assert extract_pages(data, "application/octet-stream", "underlag.PDF") == ["Innehåll."]
    assert extract_pages(data, "", "utan-filandelse") == ["Innehåll."]


def test_pdf_without_text_is_unreadable(make_pdf: MakeFile) -> None:
    with pytest.raises(UnreadableDocument, match="ingen läsbar text"):
        extract_pages(make_pdf(["", ""]), PDF_MIMETYPE, "inskannat.pdf")


def test_damaged_pdf_is_unreadable() -> None:
    with pytest.raises(UnreadableDocument, match="kunde inte läsas"):
        extract_pages(b"%PDF-1.7\nthis is not a real pdf", PDF_MIMETYPE, "trasig.pdf")


def test_password_protected_pdf_is_unreadable() -> None:
    pdf = FPDF()
    pdf.set_encryption(owner_password="agare", user_password="hemligt")
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    pdf.cell(text="Hemligt underlag")
    with pytest.raises(UnreadableDocument, match="lösenordsskyddad"):
        extract_pages(bytes(pdf.output()), PDF_MIMETYPE, "skyddad.pdf")


def test_docx_paragraphs_and_tables_become_pseudo_pages(make_docx: MakeFile) -> None:
    paragraphs = [f"Krav {n}: Anbudsgivaren ska uppfylla kravet. " + "Text " * 120 for n in range(1, 13)]
    data = make_docx(paragraphs, table=[["Krav", "Uppfylls"], ["Ansvarsförsäkring", "Ja"]])
    pages = extract_pages(data, DOCX_MIMETYPE, "underlag.docx")
    assert len(pages) > 1
    assert all(len(page) <= PSEUDO_PAGE_CHARS for page in pages)
    assert pages[0].startswith("Krav 1: Anbudsgivaren ska uppfylla kravet.")
    assert pages[-1].endswith("Krav | Uppfylls\nAnsvarsförsäkring | Ja")


def test_docx_detected_by_file_name(make_docx: MakeFile) -> None:
    data = make_docx(["Rubrik", "Brödtext om målning."])
    assert extract_pages(data, "application/octet-stream", "underlag.docx") == [
        "Rubrik\nBrödtext om målning."
    ]


def test_damaged_docx_is_unreadable() -> None:
    with pytest.raises(UnreadableDocument, match="Word-filen"):
        extract_pages(b"PK\x03\x04 not a zip", DOCX_MIMETYPE, "trasig.docx")


def test_text_is_split_on_form_feeds() -> None:
    data = "Sida ett\fSida två\n\fSida tre\f".encode()
    assert extract_pages(data, TEXT_MIMETYPE, "underlag.txt") == ["Sida ett", "Sida två", "Sida tre"]


def test_text_without_form_feeds_is_cut_at_paragraphs() -> None:
    paragraphs = [f"Stycke {n}. " + "ord " * 200 for n in range(10)]
    pages = extract_pages("\n\n".join(paragraphs).encode(), TEXT_MIMETYPE, "underlag.txt")
    assert len(pages) > 1
    assert all(len(page) <= PSEUDO_PAGE_CHARS for page in pages)
    assert pages[0].startswith("Stycke 0.")


def test_markdown_and_byte_order_mark() -> None:
    data = "\N{ZERO WIDTH NO-BREAK SPACE}# Rubrik\n\nText om städning.".encode()
    assert extract_pages(data, "text/markdown; charset=utf-8", "underlag.md") == [
        "# Rubrik\n\nText om städning."
    ]


def test_legacy_windows_encoded_text() -> None:
    data = "Förfrågningsunderlag för städtjänster".encode("cp1252")
    assert extract_pages(data, TEXT_MIMETYPE, "gammal.txt") == ["Förfrågningsunderlag för städtjänster"]


def test_unsupported_format_is_rejected() -> None:
    with pytest.raises(UnsupportedDocument, match="stöds inte"):
        extract_pages(b"\x89PNG\r\n", "image/png", "bild.png")
    with pytest.raises(UnsupportedDocument):
        extract_pages(b"\xd0\xcf\x11\xe0", "application/msword", "gammalt.doc")


def test_resolve_mimetype() -> None:
    assert resolve_mimetype("application/octet-stream", "Underlag.Docx", b"PK") == DOCX_MIMETYPE
    assert resolve_mimetype("text/markdown", "a.md", b"#") == TEXT_MIMETYPE
    assert resolve_mimetype("", "noname", b"%PDF-1.4") == PDF_MIMETYPE


# --- limits ---------------------------------------------------------------------------

MANY_PAGES = MAX_PAGES + 1


def _docx_with_body(make_docx: MakeFile, body: bytes, extra: dict[str, bytes] | None = None) -> bytes:
    """A real DOCX whose document body is replaced by `body`, plus optional extra members."""

    source = zipfile.ZipFile(io.BytesIO(make_docx(["Platshållare"])))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "word/document.xml":
                head, rest = data.split(b"<w:body>", 1)
                tail = rest[rest.index(b"<w:sectPr") :]
                data = head + b"<w:body>" + body + tail
            target.writestr(item, data)
        for name, data in (extra or {}).items():
            target.writestr(name, data)
    return out.getvalue()


def _paragraphs(count: int, text: str = "Anbudsgivaren ska uppfylla kravet.") -> bytes:
    return f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>".encode() * count


def test_docx_decompression_bomb_is_rejected_before_parsing(
    make_docx: MakeFile, monkeypatch: pytest.MonkeyPatch
) -> None:
    bomb = _docx_with_body(make_docx, b"<w:p/>" * 5_000_000)  # 30 MB of XML in about 50 kB
    assert len(bomb) < 200_000

    def must_not_parse(_stream: object) -> None:
        raise AssertionError("python-docx must not see the archive")

    monkeypatch.setattr(extract.docx, "Document", must_not_parse)
    with pytest.raises(DocumentTooLarge, match="packas upp"):
        extract_pages(bomb, DOCX_MIMETYPE, "bomb.docx")


def test_docx_element_count_catches_dense_xml_below_the_ratio_check(
    make_docx: MakeFile, monkeypatch: pytest.MonkeyPatch
) -> None:
    body = _paragraphs(1) + b"<w:p/>" * 90_000  # 540 kB: too small for the ratio check
    data = _docx_with_body(make_docx, body)
    assert extract_pages(data, DOCX_MIMETYPE, "underlag.docx") == ["Anbudsgivaren ska uppfylla kravet."]
    monkeypatch.setattr(extract, "DOCX_MAX_XML_ELEMENTS", 50_000)
    with pytest.raises(DocumentTooLarge, match="packas upp"):
        extract_pages(data, DOCX_MIMETYPE, "underlag.docx")


def test_docx_padded_with_paragraphs_is_too_long(
    make_docx: MakeFile, monkeypatch: pytest.MonkeyPatch
) -> None:
    data = _docx_with_body(make_docx, _paragraphs(1) + b"<w:p/>" * 10)
    monkeypatch.setattr(extract, "DOCX_MAX_BODY_ITEMS", 10)
    with pytest.raises(DocumentTooLarge, match="för långt"):
        extract_pages(data, DOCX_MIMETYPE, "underlag.docx")


@pytest.mark.parametrize(
    ("limit", "value"),
    [("DOCX_MAX_UNCOMPRESSED_BYTES", 10_000), ("DOCX_MAX_XML_BYTES", 10_000), ("DOCX_MAX_MEMBERS", 3)],
)
def test_docx_archive_limits(
    make_docx: MakeFile, monkeypatch: pytest.MonkeyPatch, limit: str, value: int
) -> None:
    data = _docx_with_body(make_docx, _paragraphs(400))
    monkeypatch.setattr(extract, limit, value)
    with pytest.raises(DocumentTooLarge, match="packas upp"):
        extract_pages(data, DOCX_MIMETYPE, "underlag.docx")


def test_docx_binary_parts_count_toward_size_but_not_toward_xml(
    make_docx: MakeFile, monkeypatch: pytest.MonkeyPatch
) -> None:
    image = b"\x89PNG\r\n\x1a\n" + os.urandom(2_000_000)  # incompressible, like a real photo
    data = _docx_with_body(make_docx, _paragraphs(2), extra={"word/media/image1.png": image})
    # The python-docx template alone has about 830 kB of XML and 19 000 elements.
    monkeypatch.setattr(extract, "DOCX_MAX_XML_BYTES", 1_000_000)
    monkeypatch.setattr(extract, "DOCX_MAX_XML_ELEMENTS", 20_000)
    text = "Anbudsgivaren ska uppfylla kravet."
    assert extract_pages(data, DOCX_MIMETYPE, "underlag.docx") == [f"{text}\n{text}"]
    monkeypatch.setattr(extract, "DOCX_MAX_UNCOMPRESSED_BYTES", 2_500_000)
    with pytest.raises(DocumentTooLarge, match="packas upp"):
        extract_pages(data, DOCX_MIMETYPE, "underlag.docx")


@pytest.mark.parametrize(
    "part",
    [
        b'<?xml version="1.0" encoding="UTF-7"?><w:document/>',
        "<?xml version='1.0' encoding='UTF-8'?>".encode("utf-32"),
        "<?xml version='1.0' encoding='IBM037'?><w:document/>".encode("cp037"),
    ],
    ids=["utf-7", "utf-32", "ebcdic"],
)
def test_docx_xml_in_an_encoding_that_hides_tags_is_unreadable(make_docx: MakeFile, part: bytes) -> None:
    data = _docx_with_body(make_docx, _paragraphs(1), extra={"customXml/item9.xml": part})
    with pytest.raises(UnreadableDocument, match="Word-filen"):
        extract_pages(data, DOCX_MIMETYPE, "underlag.docx")


def test_docx_xml_in_utf16_is_counted(make_docx: MakeFile, monkeypatch: pytest.MonkeyPatch) -> None:
    part = ('<?xml version="1.0" encoding="UTF-16"?><root>' + "<a/>" * 2_000 + "</root>").encode("utf-16")
    data = _docx_with_body(make_docx, _paragraphs(1), extra={"customXml/item9.xml": part})
    assert extract_pages(data, DOCX_MIMETYPE, "underlag.docx") == ["Anbudsgivaren ska uppfylla kravet."]
    monkeypatch.setattr(extract, "DOCX_MAX_XML_ELEMENTS", 1_500)
    with pytest.raises(DocumentTooLarge):
        extract_pages(data, DOCX_MIMETYPE, "underlag.docx")


def test_docx_with_a_corrupt_member_is_unreadable(make_docx: MakeFile) -> None:
    data = bytearray(make_docx(["Rubrik", "Brödtext om målning."]))
    start = data.index(b"word/document.xml") + len("word/document.xml")  # first local header
    data[start + 40 : start + 60] = bytes(20)  # damage the compressed body; the CRC no longer matches
    with pytest.raises(UnreadableDocument, match="Word-filen"):
        extract_pages(bytes(data), DOCX_MIMETYPE, "trasig.docx")


def test_text_with_too_many_pages_is_rejected() -> None:
    assert len(extract_pages(b"x\f" * MAX_PAGES, TEXT_MIMETYPE, "a.txt")) == MAX_PAGES
    with pytest.raises(DocumentTooLarge, match=f"högst {MAX_PAGES} sidor"):
        extract_pages(b"x\f" * MANY_PAGES, TEXT_MIMETYPE, "a.txt")


def test_pdf_page_count_is_checked_before_extracting_text(
    make_pdf: MakeFile, monkeypatch: pytest.MonkeyPatch
) -> None:
    data = make_pdf(["Sida ett.", "Sida två.", "Sida tre."])
    monkeypatch.setattr(extract, "MAX_PAGES", 2)

    def must_not_extract(_page: object, *args: object, **kwargs: object) -> str:
        raise AssertionError("text must not be extracted")

    monkeypatch.setattr("pypdf.PageObject.extract_text", must_not_extract)
    with pytest.raises(DocumentTooLarge, match="högst 2 sidor"):
        extract_pages(data, PDF_MIMETYPE, "lang.pdf")


def test_too_much_text_is_rejected(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(extract, "MAX_CHARS", 20)
    assert extract_pages(b"a" * 20, TEXT_MIMETYPE, "a.txt") == ["a" * 20]
    with pytest.raises(DocumentTooLarge, match="för mycket text"):
        extract_pages(b"a" * 21, TEXT_MIMETYPE, "a.txt")


def test_text_far_beyond_the_limit_is_rejected_before_it_is_split(monkeypatch: pytest.MonkeyPatch) -> None:
    def must_not_split(*_args: object, **_kwargs: object) -> list[str]:
        raise AssertionError("the text must not be split")

    monkeypatch.setattr(extract, "paginate", must_not_split)
    with pytest.raises(DocumentTooLarge, match="för mycket text"):
        extract_pages(b"a\n\n" * (MAX_CHARS + 1), TEXT_MIMETYPE, "a.txt")


def test_docx_with_too_much_text_is_rejected(make_docx: MakeFile, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(extract, "MAX_CHARS", 100)
    data = _docx_with_body(make_docx, _paragraphs(10))  # 340 characters
    with pytest.raises(DocumentTooLarge, match="för mycket text"):
        extract_pages(data, DOCX_MIMETYPE, "underlag.docx")


# --- Word: content controls, text boxes, footnotes and nested tables ----------------------

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
FOOTNOTES_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes"
FOOTNOTES_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"


def _run(text: str) -> str:
    return f"<w:r><w:t xml:space='preserve'>{text}</w:t></w:r>"


def _docx_with_xml(*body_xml: str) -> bytes:
    """A Word file whose body holds the given WordprocessingML, as a template or a form would."""

    from docx.oxml.ns import qn
    from docx.oxml.parser import parse_xml

    document = docx.Document()
    body = document.element.body
    section = body.find(qn("w:sectPr"))
    for xml in body_xml:
        element = parse_xml(f'<w:body xmlns:w="{W_NS}" xmlns:mc="{MC_NS}">{xml}</w:body>')[0]
        body.insert(list(body).index(section), element)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _with_footnotes(data: bytes, notes: dict[int, str]) -> bytes:
    """Add a footnotes part the way Word writes it: the part, its relationship and its content type."""

    footnotes = "".join(
        f'<w:footnote w:id="{number}"><w:p>{_run(text)}</w:p></w:footnote>' for number, text in notes.items()
    )
    separators = (
        '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p>'
        "</w:footnote>"
    )
    part = (
        f'<?xml version="1.0" encoding="UTF-8"?>'
        f'<w:footnotes xmlns:w="{W_NS}">{separators}{footnotes}</w:footnotes>'
    )
    source = zipfile.ZipFile(io.BytesIO(data))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for member in source.infolist():
            content = source.read(member)
            if member.filename == "word/_rels/document.xml.rels":
                relationship = f'<Relationship Id="rIdFn" Type="{FOOTNOTES_TYPE}" Target="footnotes.xml"/>'
                content = content.replace(b"</Relationships>", f"{relationship}</Relationships>".encode())
            elif member.filename == "[Content_Types].xml":
                part_name = "/word/footnotes.xml"
                override = f'<Override PartName="{part_name}" ContentType="{FOOTNOTES_CONTENT_TYPE}"/>'
                content = content.replace(b"</Types>", (override + "</Types>").encode())
            target.writestr(member, content)
        target.writestr("word/footnotes.xml", part)
    return buffer.getvalue()


def test_extraction_version_names_this_reader() -> None:
    assert extract.EXTRACTION_VERSION == 3


EAVROP_FOOTER = (
    "Sundsvalls kommun\nUpphandling\nFiktiv upphandling\nDokument\nKrav på Leverantör UH-2026-999\n"
    "Diarienummer\nUH-2026-999\nDatum\n2026-10-07\n{page}"
)


def test_an_eavrop_export_loses_its_footer_and_is_marked(make_pdf: MakeFile) -> None:
    pages = [
        "1 Krav på Leverantör\nLeverantören ska ha F-skatt.\n" + EAVROP_FOOTER.format(page=1),
        "1.2 Försäkring\nLeverantören ska ha ansvarsförsäkring.\n" + EAVROP_FOOTER.format(page=2),
    ]
    extraction = extract.extract_document(make_pdf(pages), PDF_MIMETYPE, "Krav på Leverantör.pdf")
    assert extraction.layout == "eavrop"
    assert extraction.pages == [
        "1 Krav på Leverantör\nLeverantören ska ha F-skatt.",
        "1.2 Försäkring\nLeverantören ska ha ansvarsförsäkring.",
    ]
    plain = extract.extract_document(make_pdf(["Ett vanligt underlag."]), PDF_MIMETYPE, "u.pdf")
    assert (plain.layout, plain.pages) == (None, ["Ett vanligt underlag."])


def test_docx_content_controls_are_read_in_the_body_and_inside_paragraphs() -> None:
    """A municipal template keeps the diary number and whole clauses in content controls."""

    inline = f"<w:sdt><w:sdtContent>{_run('UH-2026-120')}</w:sdtContent></w:sdt>"
    data = _docx_with_xml(
        f"<w:p>{_run('Diarienummer: ')}{inline}</w:p>",
        f"<w:sdt><w:sdtContent><w:p>{_run('Anbudsgivaren ska ha F-skatt.')}</w:p></w:sdtContent></w:sdt>",
    )
    assert extract_pages(data, DOCX_MIMETYPE, "mall.docx") == [
        "Diarienummer: UH-2026-120\nAnbudsgivaren ska ha F-skatt."
    ]


def test_docx_text_boxes_are_read_once_and_deleted_text_is_not() -> None:
    """A drawing carries its text box twice (Choice and Fallback); a tracked deletion is not the text."""

    text_box = (
        "<w:p><w:r><mc:AlternateContent><mc:Choice Requires='wps'><w:drawing><w:txbxContent>"
        f"<w:p>{_run('Anbudet ska vara giltigt i 90 dagar.')}</w:p></w:txbxContent></w:drawing></mc:Choice>"
        f"<mc:Fallback><w:pict><w:txbxContent><w:p>{_run('Anbudet ska vara giltigt i 90 dagar.')}</w:p>"
        "</w:txbxContent></w:pict></mc:Fallback></mc:AlternateContent></w:r></w:p>"
    )
    tracked = (
        f"<w:p>{_run('Priset gäller ')}<w:del><w:r><w:delText>tolv</w:delText></w:r></w:del>"
        f"<w:ins>{_run('tjugofyra')}</w:ins>{_run(' månader.')}</w:p>"
    )
    pages = extract_pages(_docx_with_xml(text_box, tracked), DOCX_MIMETYPE, "formulär.docx")
    assert pages == ["Anbudet ska vara giltigt i 90 dagar.\nPriset gäller tjugofyra månader."]


def test_docx_nested_tables_and_tabs_keep_their_text() -> None:
    cell_with_table = (
        f"<w:p>{_run('Krav')}</w:p><w:tbl><w:tr><w:tc><w:p>{_run('Inre A')}</w:p></w:tc>"
        f"<w:tc><w:p>{_run('Inre B')}</w:p></w:tc></w:tr></w:tbl>"
    )
    table = f"<w:tbl><w:tr><w:tc>{cell_with_table}</w:tc><w:tc><w:p>{_run('Ja')}</w:p></w:tc></w:tr></w:tbl>"
    tabbed = f"<w:p>{_run('Pris')}<w:r><w:tab/></w:r>{_run('450 kr/tim')}</w:p>"
    pages = extract_pages(_docx_with_xml(table, tabbed), DOCX_MIMETYPE, "tabell.docx")
    assert pages == ["Krav Inre A | Inre B | Ja\nPris 450 kr/tim"]


def test_docx_footnotes_come_after_the_body(make_docx: MakeFile) -> None:
    data = _with_footnotes(
        make_docx(["Leverantören ska ha ansvarsförsäkring."]),
        {1: "Kravet gäller inte vid uppdrag under 50 000 kronor."},
    )
    assert extract_pages(data, DOCX_MIMETYPE, "noter.docx") == [
        "Leverantören ska ha ansvarsförsäkring.\nFotnoter\n"
        "[1] Kravet gäller inte vid uppdrag under 50 000 kronor."
    ]


# --- Excel ------------------------------------------------------------------------------


def _xlsx(sheets: dict[str, list[list[object]]]) -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    first = workbook.active
    assert first is not None
    workbook.remove(first)
    for title, rows in sheets.items():
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_xlsx_sheets_become_pages_of_rows() -> None:
    import datetime as dt

    data = _xlsx(
        {
            "Priser": [
                ["Post", "Pris per timme", "Giltig från"],
                ["Målning", 450.0, dt.datetime(2026, 11, 1)],
            ],
            "Krav": [["Krav", "Ska-krav"], ["ISO 14001", True], [None, None], ["Vite", 2.5]],
            "Tomt": [[None, None]],
        }
    )
    pages = extract_pages(data, "application/octet-stream", "Bilaga 1 Prisbilaga.xlsx")
    assert pages == [
        "Blad: Priser\nPost | Pris per timme | Giltig från\nMålning | 450 | 2026-11-01",
        "Blad: Krav\nKrav | Ska-krav\nISO 14001 | Ja\nVite | 2,5",
    ]


def test_xlsx_is_detected_by_mimetype_and_damaged_workbooks_are_unreadable() -> None:
    assert resolve_mimetype(extract.XLSX_MIMETYPE, "x", b"PK") == extract.XLSX_MIMETYPE
    assert resolve_mimetype("application/octet-stream", "makro.xlsm", b"PK") == extract.XLSX_MIMETYPE
    with pytest.raises(UnreadableDocument, match="Excel-filen"):
        extract_pages(b"PK\x03\x04 not a workbook", extract.XLSX_MIMETYPE, "trasig.xlsx")

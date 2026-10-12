"""Text extraction from procurement documents, one entry per page.

Page numbers are the backbone of citations: engines see the text with "[SIDA n]"
markers and our verifier checks quotes against the same list. PDFs therefore keep
their physical pages. Formats without pages (DOCX, XLSX, plain text) are cut into
pseudo-pages of about `PSEUDO_PAGE_CHARS` characters at paragraph boundaries, so a
citation still points to a stable, readable chunk.

The cleaning keeps line breaks (lists, headings and tables stay readable for people
and models) but repairs what PDF extraction typically breaks: ligatures, odd spaces,
soft hyphens and words hyphenated across a line break.

Word files are read from their XML rather than through python-docx's text properties,
which only see ordinary runs: a municipal template keeps the diary number, dates and
whole clauses in content controls, forms put their fields in text boxes, and footnotes
hold the fine print. A requirement in any of them must not vanish without a trace.

Uploads come from anyone who picks a persona, and the app runs as one process, so
extraction must never use memory out of proportion to the upload. DOCX and XLSX are
zip archives whose XML the parsers expand in full: a small file can unpack to
gigabytes. The archive is therefore checked before parsing (`_check_office_archive`),
and every format stops at `MAX_PAGES` pages or `MAX_CHARS` characters, which is also
more than an engine can take in one call.

`EXTRACTION_VERSION` is stored with every document. Raise it when the text a stored
file would give changes, so that the worker reads the files again (`reextract_outdated`).
"""

from __future__ import annotations

import datetime as dt
import io
import logging
import re
import unicodedata
import zipfile
from collections.abc import Iterable, Iterator
from pathlib import PurePosixPath
from typing import Any, Literal, NamedTuple

import docx
import openpyxl
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml.document import CT_Document
from docx.oxml.ns import qn
from docx.oxml.parser import parse_xml
from pypdf import PdfReader

logger = logging.getLogger(__name__)

EXTRACTION_VERSION = 3
"""Version 3: e-Avrop's page footer is stripped and the export is recognised as a layout.
Version 2: Word content controls, text boxes, footnotes and nested tables; Excel sheets."""

PSEUDO_PAGE_CHARS = 3_500

MAX_PAGES = 300
"""Most pages (physical for PDF, pseudo-pages otherwise) read from one document."""
MAX_CHARS = 1_500_000
"""Most characters of cleaned text read from one document."""

DOCX_MAX_MEMBERS = 2_000
"""Most files inside a DOCX or XLSX archive; Office writes a few dozen plus one per image."""
DOCX_MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
"""Most bytes all files in an Office archive may unpack to (python-docx keeps every part)."""
DOCX_MAX_XML_BYTES = 32 * 1024 * 1024
"""Most bytes of XML in an Office archive; far more than `MAX_CHARS` of text needs."""
DOCX_MAX_XML_ELEMENTS = 1_000_000
"""Most XML elements in an Office archive. lxml and python-docx spend 100-250 bytes per
element, so element-dense XML (empty paragraphs) is the costliest kind to parse."""
DOCX_MAX_BODY_ITEMS = 100_000
"""Most top-level paragraphs and tables read from a DOCX. 300 pages hold a few tens of
thousands at most; more is padding, such as hundreds of thousands of empty paragraphs."""
DOCX_MAX_COMPRESSION_RATIO = 100
"""Highest unpacked/packed ratio for a large archive member. Office XML packs about 5-20x."""
DOCX_RATIO_MIN_BYTES = 1024 * 1024
"""Members smaller than this skip the ratio check: they cannot do harm on their own."""
XLSX_MAX_ROWS = 200_000
"""Most rows read from a workbook, empty ones included: a sheet can declare a range far
beyond its cells, and openpyxl then yields the empty rows in between."""
_ZIP_CHUNK_BYTES = 256 * 1024

PDF_MIMETYPE = "application/pdf"
DOCX_MIMETYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIMETYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
XLSM_MIMETYPE = "application/vnd.ms-excel.sheet.macroEnabled.12"
TEXT_MIMETYPE = "text/plain"
MARKDOWN_MIMETYPE = "text/markdown"

DocumentKind = Literal["pdf", "docx", "xlsx", "text"]

_KIND_BY_MIMETYPE: dict[str, DocumentKind] = {
    PDF_MIMETYPE: "pdf",
    DOCX_MIMETYPE: "docx",
    XLSX_MIMETYPE: "xlsx",
    XLSM_MIMETYPE: "xlsx",
    TEXT_MIMETYPE: "text",
    MARKDOWN_MIMETYPE: "text",
    "text/x-markdown": "text",
}
_KIND_BY_SUFFIX: dict[str, DocumentKind] = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".xlsx": "xlsx",
    ".xlsm": "xlsx",
    ".txt": "text",
    ".md": "text",
    ".markdown": "text",
}
_MIMETYPE_BY_KIND: dict[DocumentKind, str] = {
    "pdf": PDF_MIMETYPE,
    "docx": DOCX_MIMETYPE,
    "xlsx": XLSX_MIMETYPE,
    "text": TEXT_MIMETYPE,
}
SUPPORTED_SUFFIXES = frozenset(_KIND_BY_SUFFIX)
"""File endings the extraction reads; ingestion fetches only these."""

# WordprocessingML tags, resolved once (lxml compares tag strings).
_W_P = qn("w:p")
_W_TBL = qn("w:tbl")
_W_TR = qn("w:tr")
_W_TC = qn("w:tc")
_W_T = qn("w:t")
_W_TAB = qn("w:tab")
_W_BR = qn("w:br")
_W_CR = qn("w:cr")
_W_SDT = qn("w:sdt")
_W_SDT_CONTENT = qn("w:sdtContent")
_W_DEL = qn("w:del")
_W_NO_BREAK_HYPHEN = qn("w:noBreakHyphen")
_W_FOOTNOTE = qn("w:footnote")
_W_TYPE = qn("w:type")
_W_ID = qn("w:id")
# Markup compatibility: a drawing or text box comes as an AlternateContent with a Choice and
# a Fallback that hold the same text twice; only the Choice is read.
_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
_BLOCK_TAGS = frozenset({_W_P, _W_TBL})
_CONTAINER_TAGS = frozenset({_W_SDT, _W_SDT_CONTENT})
_SKIPPED_SUBTREES = frozenset({_MC_FALLBACK, _W_DEL, qn("w:moveFrom")})
# lxml ships no type stubs, so the XML elements the Word reader walks are untyped.
XmlElement = Any

_LIGATURES = str.maketrans(
    {
        "\N{LATIN SMALL LIGATURE FF}": "ff",
        "\N{LATIN SMALL LIGATURE FI}": "fi",
        "\N{LATIN SMALL LIGATURE FL}": "fl",
        "\N{LATIN SMALL LIGATURE FFI}": "ffi",
        "\N{LATIN SMALL LIGATURE FFL}": "ffl",
        "\N{LATIN SMALL LIGATURE LONG S T}": "st",
        "\N{LATIN SMALL LIGATURE ST}": "st",
    }
)
_SPACES = re.compile(
    "[\t\N{NO-BREAK SPACE}\N{FIGURE SPACE}\N{THIN SPACE}\N{HAIR SPACE}"
    "\N{NARROW NO-BREAK SPACE}\N{MEDIUM MATHEMATICAL SPACE}\N{IDEOGRAPHIC SPACE}]"
)
_INVISIBLE = re.compile(
    "[\N{ZERO WIDTH SPACE}\N{ZERO WIDTH NON-JOINER}\N{WORD JOINER}\N{ZERO WIDTH NO-BREAK SPACE}]"
)
_SOFT_HYPHEN_BREAK = re.compile("\N{SOFT HYPHEN}[ ]*\n[ ]*")
_SOFT_HYPHEN = "\N{SOFT HYPHEN}"
_HYPHEN_BREAK = re.compile(r"(?P<head>\w+)-[ ]*\n[ ]*(?P<tail>\w+)")
_RUNS_OF_SPACES = re.compile(r" {2,}")
_BLANK_LINES = re.compile(r"\n{3,}")

# A hyphen before these words marks a shared compound ("bygg- och anläggningsarbeten").
_SUSPENDED_COMPOUND_WORDS = frozenset({"och", "eller", "samt", "resp", "respektive", "som", "till"})


class DocumentError(ValueError):
    """A document cannot be used. The message is Swedish and safe to show to users."""


class UnsupportedDocument(DocumentError):
    """The file format is not supported."""


class UnreadableDocument(DocumentError):
    """The file has a supported format but no usable text could be read from it."""


class DocumentTooLarge(DocumentError):
    """The document is beyond what we extract: too many pages, too much text or a DOCX
    that unpacks to far more than its size (a decompression bomb)."""


def _too_many_pages() -> DocumentTooLarge:
    return DocumentTooLarge(
        f"Dokumentet är för långt. Anbudshjälpen kan läsa högst {MAX_PAGES} sidor. "
        "Dela upp det i mindre delar och försök igen."
    )


def _too_much_text() -> DocumentTooLarge:
    limit = f"{MAX_CHARS:,}".replace(",", " ")  # Swedish digit grouping: 1 500 000
    return DocumentTooLarge(
        f"Dokumentet innehåller för mycket text. Anbudshjälpen kan läsa högst {limit} tecken. "
        "Dela upp det i mindre delar och försök igen."
    )


def _office_too_large(kind: DocumentKind) -> DocumentTooLarge:
    name = "Excel-filen" if kind == "xlsx" else "Word-filen"
    return DocumentTooLarge(
        f"{name} blir för stor när den packas upp och kan inte läsas. "
        "Spara den som PDF eller dela upp den och försök igen."
    )


# --- Format detection -----------------------------------------------------------------


def _kind(mimetype: str, filename: str, data: bytes) -> DocumentKind:
    base_type = mimetype.split(";", 1)[0].strip().lower()
    if base_type in _KIND_BY_MIMETYPE:
        return _KIND_BY_MIMETYPE[base_type]
    suffix = PurePosixPath(filename.replace("\\", "/")).suffix.lower()
    if suffix in _KIND_BY_SUFFIX:
        return _KIND_BY_SUFFIX[suffix]
    if data.startswith(b"%PDF-"):
        return "pdf"
    raise UnsupportedDocument(
        "Filformatet stöds inte. Ladda upp underlaget som PDF, Word (.docx), Excel (.xlsx) eller textfil."
    )


def resolve_mimetype(mimetype: str, filename: str, data: bytes) -> str:
    """The canonical mimetype for a supported upload (browsers often send octet-stream)."""

    return _MIMETYPE_BY_KIND[_kind(mimetype, filename, data)]


# --- Cleaning -------------------------------------------------------------------------


def _join_hyphenated(match: re.Match[str]) -> str:
    """Rejoin a word split over a line break, keeping hyphens that belong to the text.

    "kvalifi-/ceringskrav" becomes one word; "IT-/system" and "e-/post" keep the
    hyphen (upper-case or very short first part); "bygg-/och" stays a shared compound.
    """

    head, tail = match["head"], match["tail"]
    if tail.lower() in _SUSPENDED_COMPOUND_WORDS:
        return f"{head}- {tail}"
    if len(head) >= 3 and head[-1].islower() and tail[0].islower():
        return head + tail
    return f"{head}-{tail}"


def clean_page_text(text: str) -> str:
    """Normalise one page of extracted text while keeping its line structure."""

    text = unicodedata.normalize("NFC", text).translate(_LIGATURES)
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\f", "\n")
    text = _INVISIBLE.sub("", _SPACES.sub(" ", text))
    text = _SOFT_HYPHEN_BREAK.sub("", text).replace(_SOFT_HYPHEN, "")
    lines = (_RUNS_OF_SPACES.sub(" ", line).strip() for line in text.split("\n"))
    text = "\n".join(lines)
    text = _HYPHEN_BREAK.sub(_join_hyphenated, text)
    return _BLANK_LINES.sub("\n\n", text).strip()


# --- Pseudo-pages ---------------------------------------------------------------------


def _split_long_block(block: str, size: int) -> Iterator[str]:
    """Cut a block longer than `size` at a sentence end or space where possible.

    Works with an offset instead of re-slicing the rest, so a huge block costs linear time.
    """

    block = block.strip()
    start = 0
    while len(block) - start > size:
        window = block[start : start + size]
        cut = max(window.rfind(". "), window.rfind("\n"))
        if cut < size // 2:
            cut = window.rfind(" ")
        if cut < size // 2:
            cut = size - 1
        yield window[: cut + 1].strip()
        start += cut + 1
        while start < len(block) and block[start].isspace():
            start += 1
    if start < len(block):
        yield block[start:]


def paginate(blocks: Iterable[str], *, separator: str, size: int = PSEUDO_PAGE_CHARS) -> list[str]:
    """Group text blocks (paragraphs) into pseudo-pages of at most about `size` characters."""

    pages: list[str] = []
    current: list[str] = []
    length = 0
    for block in blocks:
        for piece in _split_long_block(block, size):
            if current and length + len(separator) + len(piece) > size:
                pages.append(separator.join(current))
                current, length = [], 0
            length += (len(separator) if current else 0) + len(piece)
            current.append(piece)
    if current:
        pages.append(separator.join(current))
    return pages


# --- Formats --------------------------------------------------------------------------


def _pdf_pages(data: bytes) -> Iterator[str]:
    """Raw text per page, lazily, so that `extract_pages` can stop at its limits.

    The page count is checked before any text is extracted. pypdf caps decompressed
    streams itself (`pypdf.Configuration`).
    """

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise UnreadableDocument("PDF-filen är lösenordsskyddad och kan inte läsas.")
        pages = list(reader.pages)
    except DocumentError:
        raise
    except Exception as exc:  # pypdf raises many exception types for damaged files
        raise UnreadableDocument("PDF-filen kunde inte läsas. Den kan vara skadad.") from exc
    if len(pages) > MAX_PAGES:
        raise _too_many_pages()
    for page_no, page in enumerate(pages, start=1):
        try:
            yield page.extract_text()
        except Exception:  # keep page numbering intact even if one page is damaged
            logger.warning("Could not extract text from PDF page %d", page_no, exc_info=True)
            yield ""


_DOCX_UNREADABLE = "Word-filen kunde inte läsas. Den kan vara skadad."
_XLSX_UNREADABLE = "Excel-filen kunde inte läsas. Den kan vara skadad."
_UNREADABLE_BY_KIND: dict[DocumentKind, str] = {"docx": _DOCX_UNREADABLE, "xlsx": _XLSX_UNREADABLE}


def _unreadable(kind: DocumentKind) -> UnreadableDocument:
    return UnreadableDocument(_UNREADABLE_BY_KIND.get(kind, _DOCX_UNREADABLE))


# Starts of parts libxml2 would decode as UCS-4 or EBCDIC. OPC allows only UTF-8 and UTF-16.
_FOREIGN_XML_HEADS = (
    b"\x00\x00\xfe\xff",
    b"\xff\xfe\x00\x00",
    b"\x00\x00\x00<",
    b"<\x00\x00\x00",
    b"\x00\x00<\x00",
    b"\x00<\x00\x00",
    b"\x4c\x6f\xa7\x94",
)
_UTF16_HEADS = {
    b"\xff\xfe": "utf-16-le",
    b"\xfe\xff": "utf-16-be",
    b"<\x00?\x00": "utf-16-le",
    b"\x00<\x00?": "utf-16-be",
}
_XML_DECLARED_ENCODING = re.compile(r"<\?xml[^>]*?\bencoding\s*=\s*[\"']([A-Za-z0-9._:-]+)[\"']")
# Encodings in which "<" is always the byte 0x3C, so counting that byte counts the tags.
_BYTE_TRANSPARENT_ENCODINGS = re.compile(
    r"utf-?8|(?:us-)?ascii|iso-?8859-\d{1,2}|latin-?\d|(?:windows|cp)-?125\d", re.IGNORECASE
)


def _is_xml_part(head: bytes, kind: DocumentKind) -> bool:
    """Whether libxml2 could parse an archive member that starts with `head` as XML.

    Binary parts (images, fonts) cannot be parsed, so they are left to the size checks.
    An XML part must use an encoding in which every "<" is one 0x3C byte, otherwise the
    element count in `_check_office_archive` could be dodged (in UTF-7 "<" is "+ADw-").
    """

    if head.startswith(_FOREIGN_XML_HEADS):
        raise _unreadable(kind)
    for signature, codec in _UTF16_HEADS.items():
        if head.startswith(signature):
            declared = _XML_DECLARED_ENCODING.search(head[:400].decode(codec, "ignore"))
            if declared is not None and not declared[1].lower().startswith("utf-16"):
                raise _unreadable(kind)
            return True
    body = head.removeprefix(b"\xef\xbb\xbf").lstrip(b" \t\r\n")
    if body[:1] not in (b"<", b""):
        return False
    declared = _XML_DECLARED_ENCODING.match(body[:400].decode("latin-1"))
    if declared is not None and not _BYTE_TRANSPARENT_ENCODINGS.fullmatch(declared[1]):
        raise _unreadable(kind)
    return True


def _check_office_archive(data: bytes, kind: DocumentKind) -> None:
    """Reject a DOCX or XLSX that would unpack to far more than a parser should hold in memory.

    The declared sizes are checked first (cheap). They are trustworthy as upper bounds:
    `zipfile` never returns more than a member's declared size and fails the CRC check
    when a member is cut short. Then every member that could be XML is decompressed in
    chunks to count its bytes and elements, without building anything, because parts can
    be given any name and element density, not size, decides what parsing costs.
    """

    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
        members = archive.infolist()
    except Exception as exc:  # BadZipFile, or a damaged central directory
        raise _unreadable(kind) from exc
    if len(members) > DOCX_MAX_MEMBERS:
        raise _office_too_large(kind)
    if sum(member.file_size for member in members) > DOCX_MAX_UNCOMPRESSED_BYTES:
        raise _office_too_large(kind)
    for member in members:
        ratio = member.file_size / max(member.compress_size, 1)
        if member.file_size >= DOCX_RATIO_MIN_BYTES and ratio > DOCX_MAX_COMPRESSION_RATIO:
            raise _office_too_large(kind)
    xml_bytes = 0
    xml_elements = 0
    try:
        for member in members:
            if member.is_dir():
                continue
            with archive.open(member) as stream:
                chunk = stream.read(_ZIP_CHUNK_BYTES)
                if not _is_xml_part(chunk, kind):
                    continue
                previous = b""
                while chunk:
                    xml_bytes += len(chunk)
                    # Start tags: every "<" except those of end tags ("</"), also across chunks.
                    # In UTF-16 end tags count too, which errs on the safe side.
                    xml_elements += chunk.count(b"<") - (previous[-1:] + chunk).count(b"</")
                    if xml_bytes > DOCX_MAX_XML_BYTES or xml_elements > DOCX_MAX_XML_ELEMENTS:
                        raise _office_too_large(kind)
                    previous, chunk = chunk, stream.read(_ZIP_CHUNK_BYTES)
    except DocumentError:
        raise
    except Exception as exc:  # bad CRC, encrypted member, unsupported compression
        raise _unreadable(kind) from exc


# --- Word -----------------------------------------------------------------------------


def _children(parent: XmlElement, tags: frozenset[str]) -> Iterator[XmlElement]:
    """Children of `parent` with one of `tags`, in order, looking inside content controls.

    A content control (`w:sdt`) wraps what it holds in `w:sdtContent`: blocks in the body,
    rows in a table, cells in a row. A template or a form keeps whole sections in them.
    """

    for child in parent.iterchildren():
        if child.tag in tags:
            yield child
        elif child.tag in _CONTAINER_TAGS:
            yield from _children(child, tags)


def _inline_text(element: XmlElement) -> str:
    """The text under a paragraph or run-level element, with tabs and line breaks kept.

    The walk covers content controls, hyperlinks, fields and the text boxes of a drawing,
    which python-docx's `Paragraph.text` leaves out. It skips what is not the document's
    text: the deleted side of a tracked change and the fallback copy of a drawing. A text
    box's own paragraphs and tables become lines of their own.
    """

    parts: list[str] = []
    for child in element.iterchildren():
        tag = child.tag
        if tag in _SKIPPED_SUBTREES:
            continue
        if tag == _W_T:
            parts.append(child.text or "")
        elif tag == _W_TAB:
            parts.append("\t")
        elif tag in (_W_BR, _W_CR):
            parts.append("\n")
        elif tag == _W_NO_BREAK_HYPHEN:
            parts.append("-")
        elif tag == _W_P:
            parts.append(f"\n{_inline_text(child)}\n")
        elif tag == _W_TBL:
            parts.append("\n" + "\n".join(_table_rows(child)) + "\n")
        elif isinstance(tag, str):  # comments and processing instructions have no text
            parts.append(_inline_text(child))
    return "".join(parts)


def _block_text(block: XmlElement) -> str:
    """A paragraph as one block, a table as one line per row."""

    if block.tag == _W_TBL:
        return "\n".join(_table_rows(block))
    return _inline_text(block).strip()


def _table_rows(table: XmlElement) -> list[str]:
    """One line per row, cells separated by " | "; a nested table's rows join their cell."""

    rows: list[str] = []
    for row in _children(table, frozenset({_W_TR})):
        cells: list[str] = []
        for cell in _children(row, frozenset({_W_TC})):
            blocks = (_block_text(block) for block in _children(cell, _BLOCK_TAGS))
            text = " ".join(" ".join(block.split()) for block in blocks if block.strip())
            if text and (not cells or cells[-1] != text):  # merged cells repeat their text
                cells.append(text)
        if cells:
            rows.append(" | ".join(cells))
    return rows


def _footnotes(document: Any) -> list[str]:
    """The document's footnotes, numbered as Word numbers them, for the end of the text.

    Fine print such as "Kravet gäller inte vid …" often stands there. python-docx has no
    footnote API, so the part is read as XML; the separator notes Word always writes are
    left out, and a part that cannot be read costs the footnotes, never the document.
    """

    try:
        part = document.part.part_related_by(RELATIONSHIP_TYPE.FOOTNOTES)
        root = getattr(part, "element", None)
        if root is None:
            root = parse_xml(part.blob)  # our own archive check ran first
    except Exception:  # no footnotes part, or one Word would not have written
        return []
    notes: list[str] = []
    for note in root.iterchildren(_W_FOOTNOTE):
        if note.get(_W_TYPE):  # separator, continuationSeparator
            continue
        blocks = (_block_text(block) for block in _children(note, _BLOCK_TAGS))
        text = " ".join(" ".join(block.split()) for block in blocks if block.strip())
        if text:
            number = note.get(_W_ID)
            notes.append(f"[{number}] {text}" if number else text)
    return notes


def _docx_blocks(data: bytes) -> list[str]:
    """Text of the body's paragraphs and table rows in order, then the footnotes.

    The body is walked lazily instead of with `Document.iter_inner_content()`, which first
    builds a list of every paragraph: with many empty paragraphs that list alone costs
    more memory than the document tree.
    """

    _check_office_archive(data, "docx")
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises zip, XML and key errors for bad files
        raise UnreadableDocument(_DOCX_UNREADABLE) from exc
    root = document.element
    body = root.body if isinstance(root, CT_Document) else None
    if body is None:
        return []
    blocks: list[str] = []
    chars = 0
    for count, element in enumerate(_children(body, _BLOCK_TAGS), start=1):
        if count > DOCX_MAX_BODY_ITEMS:
            raise _too_many_pages()
        for text in _block_text(element).split("\n"):
            if text.strip():
                blocks.append(text.strip())
                chars += len(text)
        if chars > 2 * MAX_CHARS:  # cleaning shortens text a little; the exact check comes later
            raise _too_much_text()
    notes = _footnotes(document)
    if notes:
        blocks.append("Fotnoter")
        blocks.extend(notes)
    return blocks


# --- Excel ----------------------------------------------------------------------------


def _cell_value(value: object) -> str:
    """A cell as Swedish readers write it: whole numbers without ".0", decimals with a comma."""

    if value is None:
        return ""
    if isinstance(value, bool):
        return "Ja" if value else "Nej"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value).replace(".", ",")
    if isinstance(value, dt.datetime):
        if value.time() == dt.time.min:
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, dt.date):
        return value.isoformat()
    return str(value)


def _xlsx_sheets(data: bytes) -> list[list[str]]:
    """Each sheet with rows as "cell | cell" lines, the sheet's name first; empty sheets left out.

    Price lists and requirement matrices come as workbooks. Cached values are read, so a
    formula without a saved result is empty, which the citation check then cannot verify;
    that is the safe direction. Every sheet starts a new pseudo-page.
    """

    _check_office_archive(data, "xlsx")
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # not a workbook, or one openpyxl cannot read
        raise _unreadable("xlsx") from exc
    sheets: list[list[str]] = []
    rows_seen = chars = 0
    try:
        for sheet in workbook.worksheets:
            blocks = [f"Blad: {sheet.title}"]
            for row in sheet.iter_rows(values_only=True):
                rows_seen += 1
                if rows_seen > XLSX_MAX_ROWS:
                    raise _too_many_pages()
                cells = [" ".join(text.split()) for value in row if (text := _cell_value(value)).strip()]
                if cells:
                    line = " | ".join(cells)
                    blocks.append(line)
                    chars += len(line)
                    if chars > 2 * MAX_CHARS:
                        raise _too_much_text()
            if len(blocks) > 1:
                sheets.append(blocks)
    except DocumentError:
        raise
    except Exception as exc:  # damaged sheet XML
        raise _unreadable("xlsx") from exc
    finally:
        workbook.close()
    return sheets


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise UnreadableDocument("Textfilen kunde inte läsas. Spara den som UTF-8 och försök igen.")


def _text_pages(data: bytes) -> list[str]:
    text = _decode_text(data).replace("\r\n", "\n").replace("\r", "\n")
    if len(text) > 2 * MAX_CHARS:  # before splitting, which makes one object per piece
        raise _too_much_text()
    if "\f" in text:
        return [page for page in text.split("\f") if page.strip()]
    paragraphs = (paragraph.strip() for paragraph in re.split(r"\n\s*\n", text))
    return paginate((p for p in paragraphs if p), separator="\n\n")


class Extraction(NamedTuple):
    """What a file gave: its pages, and the layout they were recognised as (see `eavrop_export`)."""

    pages: list[str]
    layout: str | None


def extract_document(data: bytes, mimetype: str, filename: str) -> Extraction:
    """`extract_pages` plus the layout: an e-Avrop export loses its page footer and is marked."""

    from anbudsklar.documents.eavrop_export import detect_layout, strip_footers

    pages = _extract_pages(data, mimetype, filename)
    layout = detect_layout(pages)
    return Extraction(strip_footers(pages) if layout else pages, layout)


def extract_pages(data: bytes, mimetype: str, filename: str) -> list[str]:
    """The pages of `extract_document`, for callers that do not keep the layout."""

    return extract_document(data, mimetype, filename).pages


def _extract_pages(data: bytes, mimetype: str, filename: str) -> list[str]:
    """Extract cleaned text per page (PDF) or pseudo-page (DOCX, XLSX, text, markdown).

    A page without text stays as an empty string so that the numbering holds; callers
    count them (`empty_pages`), and the engines are told. Raises `UnsupportedDocument` for
    other formats, `UnreadableDocument` when a file is damaged, encrypted or has no text
    layer at all (scanned documents are not supported), and `DocumentTooLarge` past
    `MAX_PAGES` pages, `MAX_CHARS` characters or an Office archive limit.
    """

    kind = _kind(mimetype, filename, data)
    raw_pages: Iterable[str]
    if kind == "pdf":
        raw_pages = _pdf_pages(data)
    elif kind == "docx":
        raw_pages = paginate(_docx_blocks(data), separator="\n")
    elif kind == "xlsx":
        raw_pages = [page for sheet in _xlsx_sheets(data) for page in paginate(sheet, separator="\n")]
    else:
        raw_pages = _text_pages(data)
    pages: list[str] = []
    chars = 0
    for raw_page in raw_pages:
        if len(pages) == MAX_PAGES:
            raise _too_many_pages()
        page = clean_page_text(raw_page)
        chars += len(page)
        if chars > MAX_CHARS:
            raise _too_much_text()
        pages.append(page)
    if not any(pages):
        raise UnreadableDocument(
            "Dokumentet innehåller ingen läsbar text. Inskannade dokument utan textlager stöds inte."
        )
    return pages
